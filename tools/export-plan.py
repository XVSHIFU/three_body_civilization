"""Generate the review copy from Markdown. Optional dependency: python-docx==1.2.0."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
source = root / 'docs/PLAN-v0.3.md'
target = source.with_suffix('.docx')
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Cm(1.9)
section.left_margin = section.right_margin = Cm(1.8)
section.page_width, section.page_height = Cm(21), Cm(29.7)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)

def plain(value):
    value = re.sub(r'\[([^]]+)\]\(([^)]+)\)', r'\1（\2）', value)
    return value.replace('**', '').replace('`', '')

lines = source.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [plain(x.strip()) for x in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+', x) for x in cells):
                rows.append(cells)
            i += 1
        table = doc.add_table(rows=1, cols=len(rows[0]))
        table.style = 'Light Shading Accent 1'
        for cell, text in zip(table.rows[0].cells, rows[0]): cell.text = text
        header = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(header)
        for row in rows[1:]:
            for cell, text in zip(table.add_row().cells, row): cell.text = text
        for row in table.rows:
            no_split = OxmlElement('w:cantSplit')
            row._tr.get_or_add_trPr().append(no_split)
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs: run.font.size = Pt(9)
        doc.add_paragraph()
        continue
    if line.startswith('# '): doc.add_heading(plain(line[2:]), 0)
    elif line.startswith('## '): doc.add_heading(plain(line[3:]), 1)
    elif line.startswith('### '): doc.add_heading(plain(line[4:]), 2)
    elif line.startswith('- '): doc.add_paragraph(plain(line[2:]), 'List Bullet')
    elif re.match(r'^\d+\. ', line): doc.add_paragraph(plain(re.sub(r'^\d+\. ', '', line)), 'List Number')
    elif line: doc.add_paragraph(plain(line))
    i += 1
footer = section.footer.paragraphs[0]
footer.add_run('v0.3 · 待审阅草案 · 内容源：docs/PLAN-v0.3.md    ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = '三体文明：品质样板与交接开发计划 v0.3'
doc.core_properties.author = 'Project planning'
doc.save(target)
print(target.relative_to(root))
